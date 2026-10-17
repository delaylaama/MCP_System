"""Script to generate ~20 diverse evaluation fixtures for local AI search testing."""

import os
from pathlib import Path
from datetime import datetime, timezone, timedelta
from PIL import Image, ImageDraw

FIXTURES_DIR = Path(__file__).parent.parent / "tests" / "fixtures"
FIXTURES_DIR.mkdir(parents=True, exist_ok=True)


def create_pdf(file_path: Path, title: str, body: str, mtime: datetime):
    import fitz  # PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), f"{title}\n\n{body}", fontsize=12)
    doc.save(file_path)
    doc.close()
    os.utime(file_path, (mtime.timestamp(), mtime.timestamp()))


def create_docx(file_path: Path, title: str, body: str, mtime: datetime):
    from docx import Document
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    doc.save(file_path)
    os.utime(file_path, (mtime.timestamp(), mtime.timestamp()))


def create_text(file_path: Path, content: str, mtime: datetime):
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)
    os.utime(file_path, (mtime.timestamp(), mtime.timestamp()))


def create_leaf_image(file_path: Path, leaf_type: str, color: tuple, mtime: datetime):
    img = Image.new("RGB", (256, 256), color=(240, 248, 240))
    draw = ImageDraw.Draw(img)
    # Draw leaf shape
    draw.polygon([(128, 30), (190, 110), (170, 190), (128, 220), (86, 190), (66, 110)], fill=color)
    # Stem & veins
    draw.line([(128, 220), (128, 250)], fill=(101, 67, 33), width=5)
    draw.line([(128, 50), (128, 220)], fill=(200, 230, 200), width=2)
    draw.line([(128, 100), (160, 80)], fill=(200, 230, 200), width=2)
    draw.line([(128, 140), (155, 125)], fill=(200, 230, 200), width=2)
    draw.line([(128, 100), (96, 80)], fill=(200, 230, 200), width=2)
    draw.line([(128, 140), (101, 125)], fill=(200, 230, 200), width=2)
    img.save(file_path)
    os.utime(file_path, (mtime.timestamp(), mtime.timestamp()))


def create_non_leaf_image(file_path: Path, shape: str, color: tuple, mtime: datetime):
    img = Image.new("RGB", (256, 256), color=(245, 245, 245))
    draw = ImageDraw.Draw(img)
    if shape == "car":
        # Draw car body
        draw.rectangle([(40, 130), (216, 180)], fill=color)
        draw.polygon([(70, 130), (95, 90), (160, 90), (185, 130)], fill=color)
        # Wheels
        draw.ellipse([(65, 170), (100, 205)], fill=(20, 20, 20))
        draw.ellipse([(155, 170), (190, 205)], fill=(20, 20, 20))
    elif shape == "chart":
        # Bar chart
        draw.rectangle([(40, 120), (70, 210)], fill=(59, 130, 246))
        draw.rectangle([(90, 80), (120, 210)], fill=(16, 185, 129))
        draw.rectangle([(140, 50), (170, 210)], fill=(245, 158, 11))
        draw.rectangle([(190, 100), (220, 210)], fill=(239, 68, 68))
        draw.line([(30, 210), (230, 210)], fill=(50, 50, 50), width=2)
    elif shape == "building":
        draw.rectangle([(70, 40), (186, 230)], fill=color)
        for y in range(60, 210, 30):
            for x in range(90, 170, 25):
                draw.rectangle([(x, y), (x + 15, y + 15)], fill=(255, 255, 200))
    elif shape == "dog":
        # Dog body and head
        draw.ellipse([(60, 110), (180, 190)], fill=color)
        draw.ellipse([(160, 80), (220, 140)], fill=color)
        draw.polygon([(170, 75), (185, 50), (200, 80)], fill=(180, 130, 70))
    img.save(file_path)
    os.utime(file_path, (mtime.timestamp(), mtime.timestamp()))


def generate_all_fixtures():
    now = datetime.now(timezone.utc)
    t_2022 = now - timedelta(days=1200)
    t_2024 = now - timedelta(days=500)
    t_2025 = now - timedelta(days=120)
    t_2026_latest = now - timedelta(days=2)
    t_yesterday = now - timedelta(days=1)

    print(f"Creating evaluation fixtures in: {FIXTURES_DIR}")

    # 1. Resumes (different dates for testing 'latest' sorting)
    create_docx(
        FIXTURES_DIR / "resume_2022_thejas.docx",
        "Thejas Junior Developer Resume 2022",
        "Experienced in Python programming, HTML, CSS, JavaScript, and basic database operations. Education: BS Computer Science.",
        t_2022,
    )
    create_pdf(
        FIXTURES_DIR / "resume_2024_senior_engineer.pdf",
        "Senior Software Engineer Resume 2024",
        "Lead Backend Developer specializing in FastAPI, microservices, cloud integrations, and PostgreSQL database architecture.",
        t_2024,
    )
    create_pdf(
        FIXTURES_DIR / "resume_latest_2026_ai_lead.pdf",
        "Staff AI & Search Systems Engineer Resume (Latest 2026)",
        "Principal AI Engineer architecting local vector search, LanceDB embedding pipelines, PyTorch CLIP multimodal engines, and MCP tool protocols.",
        t_2026_latest,
    )
    create_text(
        FIXTURES_DIR / "cv_academic_researcher.txt",
        "Curriculum Vitae - Academic Researcher\nPublications in neural information retrieval, reciprocal rank fusion, and multimodal machine learning models.",
        t_2025,
    )

    # 2. Leaf and Nature Images (Testing visual CLIP embeddings)
    create_leaf_image(
        FIXTURES_DIR / "green_oak_leaf.png",
        "oak",
        (34, 139, 34),
        t_2025,
    )
    create_leaf_image(
        FIXTURES_DIR / "autumn_maple_leaf.jpg",
        "maple",
        (205, 50, 30),
        t_2024,
    )
    create_leaf_image(
        FIXTURES_DIR / "tropical_palm_leaves.png",
        "palm",
        (46, 139, 87),
        t_yesterday,
    )
    create_leaf_image(
        FIXTURES_DIR / "fern_frond_macro.png",
        "fern",
        (60, 179, 113),
        t_2025,
    )

    # 3. Non-leaf Images (Negative visual controls)
    create_non_leaf_image(
        FIXTURES_DIR / "sports_car_red.png",
        "car",
        (220, 20, 60),
        t_2024,
    )
    create_non_leaf_image(
        FIXTURES_DIR / "financial_chart_q3.png",
        "chart",
        (100, 100, 100),
        t_2025,
    )
    create_non_leaf_image(
        FIXTURES_DIR / "modern_skyscraper_glass.png",
        "building",
        (70, 130, 180),
        t_2022,
    )
    create_non_leaf_image(
        FIXTURES_DIR / "golden_retriever_dog.jpg",
        "dog",
        (218, 165, 32),
        t_2025,
    )

    # 4. Invoices & Financial Documents
    create_pdf(
        FIXTURES_DIR / "invoice_2026_01_cloud_hosting.pdf",
        "INVOICE #INV-2026-8842",
        "Vendor: Cloud Hosting Services LLC. Amount Due: $1,450.00. Description: Dedicated GPU server clusters for local embedding inference.",
        t_2026_latest,
    )
    create_text(
        FIXTURES_DIR / "quarterly_budget_2025.csv",
        "Quarter,Department,Allocation,Spent,Remaining\nQ1,Engineering,150000,142000,8000\nQ2,Research,120000,118000,2000\nQ3,Infrastructure,95000,92000,3000\nQ4,Design,45000,43000,2000\n",
        t_2025,
    )
    create_pdf(
        FIXTURES_DIR / "taxes_w2_statement_2025.pdf",
        "Form W-2 Wage and Tax Statement 2025",
        "Employer: Tech Innovations Corp. Wages, tips, other compensation: $175,000.00. Federal income tax withheld: $38,500.00.",
        t_2025,
    )

    # 5. Contracts, Specifications & Notes
    create_docx(
        FIXTURES_DIR / "nda_consulting_agreement.docx",
        "MUTUAL NON-DISCLOSURE AGREEMENT",
        "This Mutual Non-Disclosure Agreement (the 'Agreement') is entered into between Client and Contractor to protect proprietary software assets.",
        t_2024,
    )
    create_text(
        FIXTURES_DIR / "machine_learning_notes.md",
        "# Transformer Models & Reciprocal Rank Fusion Notes\n\n- Attention mechanism in MiniLM transformers\n- RRF combining scores using k=60 denominator: 1 / (60 + rank)\n- Zero-shot visual classification with CLIP ViT-B/32 cosine similarity.",
        t_yesterday,
    )
    create_text(
        FIXTURES_DIR / "meeting_minutes_march_2026.txt",
        "Meeting Minutes - March 2026 Team Standup\nAttendees: Alex, Sarah, Thejas, Elena.\nTopics discussed: Local vector database performance, LanceDB migration, OAuth refresh tokens for cloud drives.",
        t_2026_latest,
    )
    create_text(
        FIXTURES_DIR / "project_architecture_spec.md",
        "# Project MCP Drivers Architecture Specification\n\nLocal-first cross-cloud desktop application combining FastAPI, Model Context Protocol server, and Vite React frontend with LanceDB vector storage.",
        t_yesterday,
    )
    create_text(
        FIXTURES_DIR / "docker_compose_production.txt",
        "version: '3.8'\nservices:\n  backend:\n    image: mcp-drivers-backend:latest\n    ports:\n      - '8000:8000'\n    environment:\n      - HF_HUB_OFFLINE=1\n      - TRANSFORMERS_OFFLINE=1\n",
        t_2025,
    )

    count = len(list(FIXTURES_DIR.glob("*.*")))
    print(f"Successfully generated {count} test fixture files in {FIXTURES_DIR}")


if __name__ == "__main__":
    generate_all_fixtures()
